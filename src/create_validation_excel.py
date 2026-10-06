from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


BASE_DIR = Path(__file__).resolve().parent.parent

INPUT_FILE = (
    BASE_DIR / "data" / "processed" /
    "validation_review.csv"
)

OUTPUT_FILE = (
    BASE_DIR / "data" / "processed" /
    "validation_review.xlsx"
)


def main():

    print("=" * 60)
    print("CRIANDO PLANILHA DE VALIDAÇÃO")
    print("=" * 60)

    # --------------------------------------------------------
    # CSV -> XLSX
    # --------------------------------------------------------

    df = pd.read_csv(INPUT_FILE)

    df.to_excel(
        OUTPUT_FILE,
        index=False,
        engine="openpyxl"
    )

    wb = load_workbook(OUTPUT_FILE)
    ws = wb.active
    ws.title = "Validação"


    # --------------------------------------------------------
    # CABEÇALHO
    # --------------------------------------------------------

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


    # --------------------------------------------------------
    # LOCALIZAR COLUNAS
    # --------------------------------------------------------

    headers = {
        cell.value: cell.column
        for cell in ws[1]
    }


    # --------------------------------------------------------
    # LARGURAS
    # --------------------------------------------------------

    widths = {
        "validation_id": 12,
        "title": 55,
        "description": 70,
        "qwen_main_candidate": 22,
        "candidate": 22,
        "qwen_sentiment": 18,
        "qwen_stance": 18,
        "qwen_topic": 30,
        "qwen_confidence": 16,
        "human_candidate_correct": 24,
        "human_sentiment": 20,
        "human_stance": 20,
        "human_main_correct": 20,
        "human_notes": 45,
    }

    for name, width in widths.items():

        if name in headers:

            letter = ws.cell(
                row=1,
                column=headers[name]
            ).column_letter

            ws.column_dimensions[letter].width = width


    # --------------------------------------------------------
    # QUEBRA DE TEXTO
    # --------------------------------------------------------

    for row in ws.iter_rows(min_row=2):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


    # --------------------------------------------------------
    # MENUS: SIM / NAO
    # --------------------------------------------------------

    yes_no = DataValidation(
        type="list",
        formula1='"sim,nao"',
        allow_blank=True
    )

    ws.add_data_validation(yes_no)

    for column_name in [
        "human_candidate_correct",
        "human_main_correct"
    ]:

        column = ws.cell(
            row=1,
            column=headers[column_name]
        ).column_letter

        yes_no.add(
            f"{column}2:{column}{ws.max_row}"
        )


    # --------------------------------------------------------
    # MENU DE SENTIMENTO
    # --------------------------------------------------------

    sentiment = DataValidation(
        type="list",
        formula1=(
            '"positivo,negativo,neutro,'
            'misto,indeterminado"'
        ),
        allow_blank=True
    )

    ws.add_data_validation(sentiment)

    column = ws.cell(
        row=1,
        column=headers["human_sentiment"]
    ).column_letter

    sentiment.add(
        f"{column}2:{column}{ws.max_row}"
    )


    # --------------------------------------------------------
    # MENU DE STANCE
    # --------------------------------------------------------

    stance = DataValidation(
        type="list",
        formula1=(
            '"favoravel,contrario,neutro,'
            'misto,indeterminado"'
        ),
        allow_blank=True
    )

    ws.add_data_validation(stance)

    column = ws.cell(
        row=1,
        column=headers["human_stance"]
    ).column_letter

    stance.add(
        f"{column}2:{column}{ws.max_row}"
    )


    # --------------------------------------------------------
    # SALVAR
    # --------------------------------------------------------

    wb.save(OUTPUT_FILE)

    print()
    print("Registros:", len(df))
    print("Arquivo criado:")
    print(OUTPUT_FILE)

    print()
    print("Planilha pronta para revisão.")


if __name__ == "__main__":
    main()